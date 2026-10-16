"""
Document Processing Pipeline
Handles: PDF, DOCX, PPTX, TXT, Images (with OCR)
Extracts raw text, preserves page structure, stores to DB.
"""
import os
import io
import uuid
import logging
from datetime import datetime
from pathlib import Path
from typing import Optional, Tuple

from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select

from app.database.models import Document, DocumentStatus, DocumentType
from app.core.config import settings

logger = logging.getLogger(__name__)


# ─── Text Extractors ──────────────────────────────────────────────────────────

def _extract_pdf(file_bytes: bytes) -> Tuple[str, int]:
    """Extract text from PDF, returns (text, page_count)."""
    try:
        import pdfplumber
        text_parts = []
        with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
            page_count = len(pdf.pages)
            for i, page in enumerate(pdf.pages):
                t = page.extract_text()
                if t:
                    text_parts.append(f"[Page {i+1}]\n{t}")
        return "\n\n".join(text_parts), page_count
    except Exception as e:
        logger.warning(f"pdfplumber failed ({e}), trying PyPDF2")
        try:
            import PyPDF2
            reader = PyPDF2.PdfReader(io.BytesIO(file_bytes))
            page_count = len(reader.pages)
            text_parts = []
            for i, page in enumerate(reader.pages):
                t = page.extract_text() or ""
                text_parts.append(f"[Page {i+1}]\n{t}")
            return "\n\n".join(text_parts), page_count
        except Exception as e2:
            raise RuntimeError(f"PDF extraction failed: {e2}")


def _extract_docx(file_bytes: bytes) -> str:
    from docx import Document as DocxDoc
    doc = DocxDoc(io.BytesIO(file_bytes))
    parts = []
    for para in doc.paragraphs:
        if para.text.strip():
            parts.append(para.text)
    # Also extract table content
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                parts.append(row_text)
    return "\n".join(parts)


def _extract_pptx(file_bytes: bytes) -> Tuple[str, int]:
    from pptx import Presentation
    prs = Presentation(io.BytesIO(file_bytes))
    slide_texts = []
    for i, slide in enumerate(prs.slides):
        parts = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                parts.append(shape.text.strip())
        if parts:
            slide_texts.append(f"[Slide {i+1}]\n" + "\n".join(parts))
    return "\n\n".join(slide_texts), len(prs.slides)


def _extract_txt(file_bytes: bytes) -> str:
    for enc in ("utf-8", "utf-16", "latin-1"):
        try:
            return file_bytes.decode(enc)
        except UnicodeDecodeError:
            continue
    return file_bytes.decode("utf-8", errors="replace")


def _extract_image_ocr(file_bytes: bytes) -> str:
    """OCR via pytesseract. Falls back gracefully if Tesseract not installed."""
    try:
        import pytesseract
        from PIL import Image
        if settings.tesseract_cmd and settings.tesseract_cmd != "tesseract":
            pytesseract.pytesseract.tesseract_cmd = settings.tesseract_cmd
        img = Image.open(io.BytesIO(file_bytes))
        return pytesseract.image_to_string(img)
    except Exception as e:
        logger.warning(f"OCR failed: {e}. Returning empty string.")
        return ""


# ─── Main Processor ───────────────────────────────────────────────────────────

class DocumentProcessor:

    @staticmethod
    def detect_type(filename: str) -> DocumentType:
        ext = Path(filename).suffix.lower()
        return {
            ".pdf":  DocumentType.pdf,
            ".docx": DocumentType.docx,
            ".doc":  DocumentType.docx,
            ".pptx": DocumentType.pptx,
            ".ppt":  DocumentType.pptx,
            ".txt":  DocumentType.txt,
            ".md":   DocumentType.txt,
            ".png":  DocumentType.image,
            ".jpg":  DocumentType.image,
            ".jpeg": DocumentType.image,
            ".webp": DocumentType.image,
            ".tiff": DocumentType.image,
        }.get(ext, DocumentType.txt)

    @staticmethod
    def extract_text(file_bytes: bytes, doc_type: DocumentType) -> Tuple[str, int]:
        """Extract raw text. Returns (text, page_count)."""
        if doc_type == DocumentType.pdf:
            return _extract_pdf(file_bytes)
        elif doc_type == DocumentType.docx:
            text = _extract_docx(file_bytes)
            return text, 1
        elif doc_type == DocumentType.pptx:
            return _extract_pptx(file_bytes)
        elif doc_type == DocumentType.image:
            text = _extract_image_ocr(file_bytes)
            return text, 1
        else:
            return _extract_txt(file_bytes), 1

    @staticmethod
    async def save_file(file_bytes: bytes, filename: str, user_id: str) -> str:
        """Save uploaded file to disk, return relative path."""
        safe_name = f"{uuid.uuid4().hex}_{Path(filename).name}"
        user_dir = os.path.join(settings.upload_dir, user_id)
        os.makedirs(user_dir, exist_ok=True)
        file_path = os.path.join(user_dir, safe_name)
        with open(file_path, "wb") as f:
            f.write(file_bytes)
        return file_path

    @staticmethod
    async def process_document(
        db: AsyncSession,
        document_id: str,
        embedding_service=None,
    ) -> Document:
        """
        Full pipeline:
        1. Load file bytes from disk
        2. Extract text
        3. Update document record
        4. Chunk + embed via EmbeddingService (if provided)
        """
        stmt = select(Document).where(Document.id == document_id)
        result = await db.execute(stmt)
        doc = result.scalar_one_or_none()

        if not doc:
            raise ValueError(f"Document {document_id} not found")

        # Mark processing
        doc.status = DocumentStatus.processing
        await db.commit()

        try:
            with open(doc.file_path, "rb") as f:
                file_bytes = f.read()

            text, page_count = DocumentProcessor.extract_text(file_bytes, doc.file_type)
            doc.raw_text   = text.strip()
            doc.page_count = page_count
            doc.status     = DocumentStatus.ready
            doc.processed_at = datetime.utcnow()

            await db.commit()
            await db.refresh(doc)

            # Chunk and embed
            if embedding_service and doc.raw_text:
                chunk_count = await embedding_service.embed_document(db, doc)
                doc.chunk_count = chunk_count
                await db.commit()

            logger.info(f"Document {doc.file_name} processed: {doc.chunk_count} chunks")
            return doc

        except Exception as e:
            doc.status        = DocumentStatus.failed
            doc.error_message = str(e)
            await db.commit()
            logger.error(f"Document processing failed: {e}", exc_info=True)
            raise
